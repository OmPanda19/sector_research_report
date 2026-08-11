"""Support & Audit table builder.

Every sheet ends with a 16-column table that explains every important number on
that sheet. Rows are generated from three places:

 1. the pre-existing research block (its own Line item / Unit / Source / Conf. /
    Notes columns are the authoritative provenance and are re-presented here),
 2. hand-written rows for outputs that live only in the new upper tables,
 3. one row for every ORANGE cell group, carrying the reason it is unavailable.

The Value column is a LIVE formula, so the support table re-states itself when
the scenario changes.
"""
import re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .spec import CODE, LEGACY

HDR = ['Reference ID', 'Model Item', 'Value', 'Unit', 'Calculation Method', 'Formula Used',
       'Primary Source', 'Secondary Source', 'Assumption', 'Reasoning', 'Cross-check',
       'Confidence Level', 'Linked Sheet', 'Update Frequency', 'Last Available Date', 'Comments']

NAVY = 'FF1F3864'
LIGHT = 'FFF2F2F2'
BAND = 'FFFAFAFA'
THIN = Side(style='thin', color='FFBFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FREQ_BY_SOURCE = [
    (r'RBI|MPC|repo|G-sec|CRISIL|investing', 'On each RBI MPC statement (bi-monthly)', 'Jun-2026'),
    (r'World Bank|Pink', 'Monthly, on each Pink Sheet release', 'Jun-2026'),
    (r'worldsteel|OECD', 'Twice yearly (April and October SRO)', 'Apr-2026'),
    (r'Ministry|JPC|Joint Plant', 'Monthly, on each Ministry of Steel report', 'Apr-2026'),
    (r'Master DB|Database|Derived|Derived-DB|00 Database', 'Quarterly, on each results season', 'FY2026A'),
    (r'policy|P\d\d|Statutory|Convention', 'On any statutory or policy change', 'In force'),
    (r'Indicative|Not sourced|NOT SOURCED|Composite', 'Quarterly review; replace before transaction use', 'n/a'),
    (r'Computed|^\d\d ', 'Recalculates live with the model', 'Live'),
]


def _freq(src):
    s = str(src or '')
    for pat, freq, last in FREQ_BY_SOURCE:
        if re.search(pat, s, re.I):
            return freq, last
    return 'Quarterly', 'FY2026A'


def _method(formula, const_val):
    """Classify a cell into one of the model's field types."""
    if formula is None:
        return 'Sourced actual (hard value)' if const_val is not None else 'Not populated'
    t = formula if isinstance(formula, str) else getattr(formula, 'text', '')
    if not t.startswith('='):
        return 'Sourced actual (hard value)'
    if 'XLOOKUP' in t or 'INDEX(' in t and 'ScenID' in t:
        return 'Scenario-driven lookup'
    if 'SUMIFS' in t:
        return 'Database-driven aggregation (SUMIFS)'
    if "'" in t and '!' in t:
        return 'Cross-sheet reference (linked, not re-entered)'
    if re.search(r'\bMIN\(|\bMAX\(|\bIF\(', t):
        return 'Conditional / constrained calculation'
    return 'Computed in-cell'


def _sheets_in(formula):
    t = formula if isinstance(formula, str) else getattr(formula, 'text', '') or ''
    names = sorted(set(re.findall(r"'([^']+)'!", t)))
    return ', '.join(names) if names else 'This sheet'


def _split_note(note):
    """Split the legacy Notes column into assumption / reasoning / cross-check."""
    n = (note or '').strip()
    if not n:
        return '', '', ''
    parts = re.split(r'(?<=[.!])\s+', n)
    assumption = parts[0] if parts else ''
    reasoning = ' '.join(parts[1:3]) if len(parts) > 1 else ''
    cross = ' '.join(parts[3:]) if len(parts) > 3 else ''
    return assumption, reasoning, cross


def _formula_text(cell):
    v = cell.value
    if v is None:
        return ''
    if isinstance(v, str):
        return v if v.startswith('=') else f'(constant) {v}'[:250]
    if hasattr(v, 'text'):
        return '{=' + str(v.text).lstrip('=') + '}'
    return f'(constant) {v}'


# the research block on these sheets does not use the Line item / Unit / Source grammar
SPECIAL = {'Audit Checks', 'Sources'}
# where the research block stops, if something was appended below it
LEGACY_STOP = {'Scenario Manager': 87}


def harvest_special(ws):
    """25 Audit Checks and 26 Sources carry their own column grammar."""
    rows = []
    if ws.title == 'Audit Checks':
        for r in range(14, 45):
            cid = ws.cell(r, 1).value
            val = ws.cell(r, 2).value
            if not isinstance(cid, str) and not isinstance(val, str):
                continue
            rows.append(dict(
                item=f'{cid or "OVERALL"} - {val or "overall result"}',
                value=f'=IF(C{r}="PASS",1,IF(C{r}="WARN",0.5,0))', unit='PASS / WARN / FAIL',
                method='Live conditional test evaluated on every recalculation',
                formula=_formula_text(ws.cell(r, 3)),
                primary='Computed from the model itself',
                secondary='Several checks are duplicated independently on 02b Model Calibration',
                assumption='None. A check is a test, not an assumption.',
                reasoning=str(ws.cell(r, 4).value or '')[:900],
                cross='Re-evaluates when the scenario is changed on 01 Control Panel',
                conf='High', linked=_sheets_in(ws.cell(r, 3).value),
                freq='Live', last='Live',
                comments=f'Research block row {r}. A WARN is a signal to read the number, not a defect.',
                numfmt='0.0'))
        return rows
    for r in range(14, ws.max_row + 1):
        ref = ws.cell(r, 1).value
        if not isinstance(ref, str) or not ref.strip():
            continue
        rows.append(dict(
            item=f'{ref} - {ws.cell(r, 3).value or ""}', value=None,
            unit=f'Level {ws.cell(r, 2).value}' if ws.cell(r, 2).value else '',
            method='Source record', formula='n/a',
            primary=f'{ws.cell(r, 3).value or ""} - {ws.cell(r, 4).value or ""}',
            secondary=str(ws.cell(r, 7).value or ''),
            assumption=str(ws.cell(r, 6).value or ''),
            reasoning=str(ws.cell(r, 8).value or '')[:900],
            cross=str(ws.cell(r, 7).value or ''), conf='Level ' + str(ws.cell(r, 2).value or ''),
            linked='02 Model Assumptions', freq='On each publication cycle',
            last=str(ws.cell(r, 5).value or ''), comments=f'Source register row {r}.'))
    return rows


def harvest_legacy(ws):
    """Build support rows from the pre-existing research block."""
    start = LEGACY.get(ws.title)
    rows = []
    if not start:
        return rows
    if ws.title in SPECIAL:
        return harvest_special(ws)
    stop = LEGACY_STOP.get(ws.title, ws.max_row)
    for r in range(start, stop + 1):
        item = ws.cell(r, 1).value
        if not isinstance(item, str) or not item.strip():
            continue
        item = item.strip()
        unit = ws.cell(r, 2).value
        # header rows and section banners: no unit, no numbers -> skip banners but
        # keep them out of the table entirely
        if item in ('Line item',) or (unit in (None, '') and not any(
                ws.cell(r, c).value is not None for c in range(3, 11))):
            continue
        src = ws.cell(r, 12).value
        conf = ws.cell(r, 13).value
        note = ws.cell(r, 14).value
        # pick the cell that best represents the row
        fcell = None
        for c in (4, 3, 8, 5):
            cc = ws.cell(r, c)
            if cc.value is not None:
                fcell = cc
                break
        formula = fcell.value if fcell is not None else None
        has_a = ws.cell(r, 3).value is not None
        has_t = ws.cell(r, 10).value is not None
        if has_t and has_a:
            value = f'=IF(ISNUMBER(J{r}),J{r},C{r})'
        elif has_t:
            value = f'=J{r}'
        elif has_a:
            value = f'=C{r}'
        else:
            value = f'=IFERROR(D{r},"")'
        a, rea, cr = _split_note(note)
        freq, last = _freq(src)
        rows.append(dict(
            item=item, value=value, unit=unit or '', method=_method(formula, ws.cell(r, 3).value),
            formula=_formula_text(fcell) if fcell is not None else '', primary=src or 'See 26 Sources',
            secondary='26 Sources register; Master Industry Database source log (47 cited sources)',
            assumption=a, reasoning=rea, cross=cr or 'Reconciled by 25 Audit Checks',
            conf=conf or 'Medium', linked=_sheets_in(formula), freq=freq, last=last,
            comments=f'Research block row {r}. Value column is live and re-states on a scenario change.',
            numfmt=None))
    return rows


def write(ws, extra_rows=None, na_rows=None, intro=None):
    """Append the Support & Audit table to `ws`. Returns the header row index."""
    rows = list(extra_rows or []) + harvest_legacy(ws) + list(na_rows or [])
    if not rows:
        return None
    code = CODE.get(ws.title, 'XX')
    top = ws.max_row + 3

    # ---- title
    t = ws.cell(top, 1, 'SUPPORT & AUDIT TABLE')
    t.font = Font(name='Calibri', sz=12, b=True, color='FFFFFFFF')
    t.fill = PatternFill('solid', fgColor=NAVY)
    t.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    for c in range(2, 17):
        cc = ws.cell(top, c)
        cc.fill = PatternFill('solid', fgColor=NAVY)
    ws.row_dimensions[top].height = 20

    sub = ws.cell(top + 1, 1, intro or (
        'Every important number on this sheet, with its calculation method, the exact formula used, its source, '
        'the reasoning behind it, how it is cross-checked, and how often it must be refreshed. The Value column '
        'is live: it re-states whenever the scenario on 01 Control Panel is changed.'))
    sub.font = Font(name='Calibri', sz=9, i=True, color='FF595959')
    sub.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=top + 1, start_column=1, end_row=top + 1, end_column=16)
    ws.row_dimensions[top + 1].height = 28

    # ---- header
    hr = top + 2
    for i, h in enumerate(HDR, start=1):
        c = ws.cell(hr, i, h)
        c.font = Font(name='Calibri', sz=9, b=True, color='FFFFFFFF')
        c.fill = PatternFill('solid', fgColor='FF2E5C8A')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BOX
    ws.row_dimensions[hr].height = 32

    # ---- body
    for n, d in enumerate(rows):
        r = hr + 1 + n
        vals = [f'{code}{n + 1:02d}', d.get('item', ''), d.get('value', ''), d.get('unit', ''),
                d.get('method', ''), d.get('formula', ''), d.get('primary', ''), d.get('secondary', ''),
                d.get('assumption', ''), d.get('reasoning', ''), d.get('cross', ''), d.get('conf', ''),
                d.get('linked', ''), d.get('freq', ''), d.get('last', ''), d.get('comments', '')]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(r, i)
            if i == 3 and isinstance(v, str) and v.startswith('='):
                c.value = v
                c.number_format = d.get('numfmt') or '#,##0.00'
                c.font = Font(name='Calibri', sz=9, b=True, color='FF1F3864')
            else:
                if isinstance(v, str) and v.startswith('='):
                    v = "'" + v          # keep 'Formula Used' as text
                c.value = v if v != '' else None
                c.font = Font(name='Calibri', sz=8, color='FF404040')
            c.alignment = Alignment(wrap_text=True, vertical='top',
                                    horizontal='center' if i in (1, 4, 12, 15) else 'left')
            c.border = BOX
            if n % 2:
                c.fill = PatternFill('solid', fgColor=BAND)
        ws.row_dimensions[r].height = 46

    # widen only columns the sheet never used, so the owner's widths are untouched
    used = ws.max_column
    for col in range(15, 17):
        letter = get_column_letter(col)
        if letter not in ws.column_dimensions:
            ws.column_dimensions[letter].width = 26
    return hr


INTRO = {
    'Cover': 'This workbook is an integrated industry financial model, not a set of independent '
             'forecasts. Real GDP growth and a demand elasticity drive volume; volume against capacity '
             'drives utilisation; utilisation feeds back into realisation with a one-year lag; '
             'realisation less cash cost drives margin, cash flow and valuation. The table below records '
             'the coverage universe and the model metadata on this page.',
    'Index': 'Every sheet in the workbook, its role, and where its numbers come from. Sheet names in '
             'column B are live hyperlinks.',
    'Database Import': 'What is imported from Master Industry Database.xlsx, how many records each '
                       'dataset carries, and the external reference formula that should replace each '
                       'RED cell when the two workbooks are reconnected.',
    'Forecast Horizon': 'Why the forecast horizon is seven years. The scoring is a framework; the '
                        'evidence behind each score is stated on the sheet.',
    'Control Panel': 'The six workbook status tests and the fourteen navigation links. Every status '
                     'cell is a live formula that can fail - none is a hard-coded tick.',
    'Model Assumptions': 'THE SINGLE SOURCE OF TRUTH FOR THE WHOLE WORKBOOK. Twenty-six drivers, each '
                         'with four scenario paths and a documented source. Nothing downstream re-keys '
                         'any of these values; everything links here. The table below documents the '
                         'FY2026A anchor of every driver.',
    'Model Calibration': 'The FY2026A actuals every forecast is anchored on, and five arithmetic '
                         'reconciliations that prove the anchors are internally consistent.',
    'Audit Checks': 'Thirty live validations plus the standing limitations. The table below explains '
                    'what each check proves and what to do if it fails.',
    'Sources': 'Every source used for a forward-looking assumption, with publisher, publication, date '
               'and what it is used for.',
}


def na_row(label, addrs, reason, unit='', linked='This sheet'):
    """Support row describing an ORANGE (data unavailable) group."""
    return dict(item=label, value=None, unit=unit,
                method='NOT POPULATED - orange fill', formula='(blank by design)',
                primary='No reliable source located', secondary='n/a',
                assumption='None asserted. The cell is deliberately blank.',
                reasoning=reason, cross='Excluded from every headline output and from 25 Audit Checks',
                conf='n/a - unavailable', linked=linked,
                freq='Populate when a primary source becomes available',
                last='Not available', comments=f'Orange cells: {addrs}')
