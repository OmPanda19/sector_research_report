"""Inspect workbook features that must survive an openpyxl round-trip."""
import sys, zipfile, collections, json
import openpyxl
from openpyxl.utils import get_column_letter

path = sys.argv[1]

def profile(p):
    wb = openpyxl.load_workbook(p)
    out = {}
    out['sheets'] = wb.sheetnames
    out['defined_names'] = sorted(wb.defined_names.keys())
    per = {}
    for ws in wb.worksheets:
        d = {}
        d['dims'] = f"{ws.max_row}x{ws.max_column}"
        d['merged'] = len(ws.merged_cells.ranges)
        d['dv'] = len(ws.data_validations.dataValidation)
        d['cf'] = len(ws.conditional_formatting._cf_rules)
        d['tables'] = len(ws.tables)
        d['col_dims'] = len(ws.column_dimensions)
        d['row_dims'] = len(ws.row_dimensions)
        d['freeze'] = ws.freeze_panes
        d['hyperlinks'] = len(ws._hyperlinks)
        d['charts'] = len(ws._charts)
        d['images'] = len(ws._images)
        nfills = 0; nfonts = 0; nborders = 0; nnumfmt = 0; ncells = 0
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    ncells += 1
                if c.has_style:
                    if c.fill is not None and c.fill.fgColor is not None and c.fill.patternType:
                        nfills += 1
                    if c.font is not None and (c.font.b or (c.font.color is not None and c.font.color.rgb not in (None, '00000000'))):
                        nfonts += 1
                    if c.border is not None and (c.border.bottom.style or c.border.top.style or c.border.left.style or c.border.right.style):
                        nborders += 1
                    if c.number_format not in ('General',):
                        nnumfmt += 1
        d.update(cells=ncells, fills=nfills, fonts=nfonts, borders=nborders, numfmt=nnumfmt)
        per[ws.title] = d
    out['per'] = per
    z = zipfile.ZipFile(p)
    out['parts'] = sorted(n for n in z.namelist())
    return out

a = profile(path)
print(json.dumps({'sheets': len(a['sheets']), 'defined_names': a['defined_names']}, indent=1))
tot = collections.Counter()
for name, d in a['per'].items():
    for k, v in d.items():
        if isinstance(v, int):
            tot[k] += v
print('TOTALS', dict(tot))
print('PARTS')
for p in a['parts']:
    print('  ', p)
print('PER SHEET')
for name, d in a['per'].items():
    print(f"  {name:32s} {d['dims']:>9s} cells={d['cells']:5d} merged={d['merged']:4d} dv={d['dv']:2d} cf={d['cf']:2d} tbl={d['tables']:2d} colw={d['col_dims']:3d} rowh={d['row_dims']:4d} fills={d['fills']:5d} fonts={d['fonts']:5d} bord={d['borders']:5d} nfmt={d['numfmt']:5d} hl={d['hyperlinks']:3d} frz={d['freeze']}")
