"""Print a rectangular window of a sheet as text, for eyeballing layout.

Run:  python tools/show.py "Sheet Name" firstrow lastrow [lastcol] [workbook.xlsx]
"""
import sys

import openpyxl
from openpyxl.utils import get_column_letter


def main(sheet, r0, r1, maxcol=12, path='Industry Financial Model.xlsx'):
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    r0, r1, maxcol = int(r0), int(r1), int(maxcol)
    widths = [22] + [16] * (maxcol - 1)
    head = 'row  ' + ''.join(get_column_letter(c + 1).ljust(widths[c])
                             for c in range(maxcol))
    print(head)
    print('-' * len(head))
    for r in range(r0, r1 + 1):
        out = f'{r:<5}'
        for c in range(1, maxcol + 1):
            cell = ws.cell(r, c)
            v = cell.value
            s = '' if v is None else str(v)
            fill = cell.fill
            tag = ''
            if fill is not None and fill.patternType == 'solid':
                rgb = getattr(fill.fgColor, 'rgb', None)
                if rgb == 'FFFFD9A0':
                    tag = '[ext]'
                elif rgb == 'FFFFC7CE':
                    tag = '[MISS]'
            s = (tag + s)[:widths[c - 1] - 1]
            out += s.ljust(widths[c - 1])
        print(out.rstrip())


if __name__ == '__main__':
    main(*sys.argv[1:])
