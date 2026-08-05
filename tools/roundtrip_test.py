"""Round-trip IFM_user.xlsx through openpyxl and diff every style attribute + value."""
import openpyxl, sys, itertools
from openpyxl.utils import get_column_letter

src = 'incoming/IFM_user.xlsx'
dst = 'incoming/IFM_roundtrip.xlsx'

wb = openpyxl.load_workbook(src)
wb.save(dst)

a = openpyxl.load_workbook(src)
b = openpyxl.load_workbook(dst)

problems = []
if a.sheetnames != b.sheetnames:
    problems.append(('SHEETNAMES', a.sheetnames, b.sheetnames))

def font_sig(f):
    return (f.name, f.sz, f.b, f.i, f.u, (f.color.rgb if f.color is not None else None),
            (f.color.theme if f.color is not None else None), (f.color.tint if f.color is not None else None))
def fill_sig(f):
    return (f.patternType, (f.fgColor.rgb if f.fgColor is not None else None),
            (f.fgColor.theme if f.fgColor is not None else None),
            (f.bgColor.rgb if f.bgColor is not None else None))
def bd_sig(b_):
    return tuple((getattr(b_, s).style, (getattr(b_, s).color.rgb if getattr(b_, s).color is not None else None))
                 for s in ('left', 'right', 'top', 'bottom', 'diagonal'))
def al_sig(al):
    return (al.horizontal, al.vertical, al.wrapText, al.indent, al.textRotation, al.shrinkToFit)

counts = dict(value=0, font=0, fill=0, border=0, align=0, numfmt=0, prot=0)
for sn in a.sheetnames:
    wa, wbs = a[sn], b[sn]
    if (wa.max_row, wa.max_column) != (wbs.max_row, wbs.max_column):
        problems.append(('DIMS', sn, (wa.max_row, wa.max_column), (wbs.max_row, wbs.max_column)))
    for r in range(1, max(wa.max_row, wbs.max_row) + 1):
        for c in range(1, max(wa.max_column, wbs.max_column) + 1):
            ca, cb = wa.cell(r, c), wbs.cell(r, c)
            va, vb = ca.value, cb.value
            if type(va).__name__ == 'ArrayFormula':
                va = ('AF', va.ref, va.text)
            if type(vb).__name__ == 'ArrayFormula':
                vb = ('AF', vb.ref, vb.text)
            if va != vb:
                counts['value'] += 1
                if counts['value'] < 8:
                    problems.append(('VALUE', sn, f'{get_column_letter(c)}{r}', repr(va)[:80], repr(vb)[:80]))
            for key, fn, attr in (('font', font_sig, 'font'), ('fill', fill_sig, 'fill'),
                                  ('border', bd_sig, 'border'), ('align', al_sig, 'alignment')):
                if fn(getattr(ca, attr)) != fn(getattr(cb, attr)):
                    counts[key] += 1
                    if counts[key] < 4:
                        problems.append((key.upper(), sn, f'{get_column_letter(c)}{r}',
                                         str(fn(getattr(ca, attr)))[:90], str(fn(getattr(cb, attr)))[:90]))
            if ca.number_format != cb.number_format:
                counts['numfmt'] += 1
                if counts['numfmt'] < 4:
                    problems.append(('NUMFMT', sn, f'{get_column_letter(c)}{r}', ca.number_format, cb.number_format))
    # structural
    for label, xa, xb in (
        ('merged', sorted(str(x) for x in wa.merged_cells.ranges), sorted(str(x) for x in wbs.merged_cells.ranges)),
        ('dv', len(wa.data_validations.dataValidation), len(wbs.data_validations.dataValidation)),
        ('cf', len(wa.conditional_formatting._cf_rules), len(wbs.conditional_formatting._cf_rules)),
        ('tables', sorted(wa.tables), sorted(wbs.tables)),
        ('colw', {k: (v.width, v.hidden) for k, v in wa.column_dimensions.items()},
                 {k: (v.width, v.hidden) for k, v in wbs.column_dimensions.items()}),
        ('rowh', {k: (v.height, v.hidden) for k, v in wa.row_dimensions.items()},
                 {k: (v.height, v.hidden) for k, v in wbs.row_dimensions.items()}),
        ('sheetview', (wa.sheet_view.showGridLines, wa.sheet_view.zoomScale, wa.freeze_panes),
                      (wbs.sheet_view.showGridLines, wbs.sheet_view.zoomScale, wbs.freeze_panes)),
        ('tabcolor', wa.sheet_properties.tabColor, wbs.sheet_properties.tabColor),
    ):
        if xa != xb:
            problems.append(('STRUCT-' + label, sn, str(xa)[:200], str(xb)[:200]))

print('COUNTS', counts)
print('PROBLEMS', len(problems))
for p in problems[:80]:
    print('  ', p)
