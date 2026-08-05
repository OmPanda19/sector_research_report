import openpyxl, os
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook('incoming/MDB_user.xlsx', data_only=False)
os.makedirs('incoming/mdb', exist_ok=True)
print('SHEETS:', wb.sheetnames)
for ws in wb.worksheets:
    lines = [f"### {ws.title}  {ws.max_row}x{ws.max_column}"]
    for r in range(1, ws.max_row + 1):
        parts = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            if isinstance(v, str):
                v = v.replace('\n', ' ')[:60]
            parts.append(f"{get_column_letter(c)}={v}")
        lines.append(f"R{r:<4d} " + (" | ".join(parts) if parts else "--blank--"))
    fn = 'incoming/mdb/' + ws.title.replace(' ', '_').replace('/', '_') + '.txt'
    open(fn, 'w').write("\n".join(lines) + "\n")
    print(f"  {ws.title:32s} {ws.max_row:4d}x{ws.max_column:3d} -> {fn}")
