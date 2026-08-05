#!/usr/bin/env python3
"""Structural verification of the delivered workbook: does it open, and is it well formed?"""
import sys, zipfile
from xml.etree import ElementTree
import openpyxl

OUT = 'Industry Financial Model.xlsx'
bad = 0

# 1. every part must be well-formed XML
z = zipfile.ZipFile(OUT)
n_xml = 0
for name in z.namelist():
    if not name.endswith(('.xml', '.rels')):
        continue
    n_xml += 1
    try:
        ElementTree.fromstring(z.read(name))
    except ElementTree.ParseError as e:
        print(f'  MALFORMED {name}: {e}')
        bad += 1
print(f'{n_xml} XML parts, all well formed' if not bad else f'{bad} malformed parts')

# 2. no dangling relationship ids
import re
names = set(z.namelist())
for rel in [n for n in names if n.endswith('.rels')]:
    base = rel.rsplit('_rels/', 1)[0]
    root = ElementTree.fromstring(z.read(rel))
    for r in root:
        tgt = r.get('Target')
        if r.get('TargetMode') == 'External' or tgt is None or tgt.startswith('http'):
            continue
        path = tgt.lstrip('/')
        if tgt.startswith('/'):
            resolved = path
        else:
            parts = (base + path).split('/')
            out = []
            for p in parts:
                if p == '..':
                    out.pop()
                elif p not in ('.', ''):
                    out.append(p)
            resolved = '/'.join(out)
        if resolved not in names:
            print(f'  DANGLING {rel} -> {tgt} ({resolved})')
            bad += 1
print('all internal relationships resolve' if not bad else 'dangling relationships found')

# 3. no external workbook links left
ext = [n for n in names if 'externalLink' in n]
print(f'external workbook links: {len(ext)} (expected 0)')
if ext:
    bad += 1

# 4. reload with openpyxl and count everything
wb = openpyxl.load_workbook(OUT)
ch = sum(len(ws._charts) for ws in wb.worksheets)
f = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            t = v if isinstance(v, str) else getattr(v, 'text', None)
            if isinstance(t, str) and t.startswith('='):
                f += 1
hl = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.hyperlink is not None)
print(f'sheets {len(wb.sheetnames)} | charts {ch} | formulas {f} | hyperlinks {hl} | '
      f'defined names {sorted(wb.defined_names)}')
print(f'fullCalcOnLoad = {wb.calculation.fullCalcOnLoad}')
if not wb.calculation.fullCalcOnLoad:
    bad += 1
if hl < 40:
    print('  too few hyperlinks'); bad += 1

# 5. round-trip: save a copy and confirm it still loads with the same counts
wb.save('incoming/_verify_roundtrip.xlsx')
wb2 = openpyxl.load_workbook('incoming/_verify_roundtrip.xlsx')
ch2 = sum(len(ws._charts) for ws in wb2.worksheets)
print(f'round-trip: sheets {len(wb2.sheetnames)} charts {ch2}')
if (len(wb2.sheetnames), ch2) != (len(wb.sheetnames), ch):
    print('  round-trip lost content')
    bad += 1

import os
print(f'file size {os.path.getsize(OUT) / 1024:.0f} KB')
print('RESULT:', 'OK' if bad == 0 else f'{bad} PROBLEM(S)')
sys.exit(1 if bad else 0)
