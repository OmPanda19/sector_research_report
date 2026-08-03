#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build "Master Industry Database.xlsx" - Indian Steel Industry.

Run:  python3 build_workbook.py
Output: ./Master Industry Database.xlsx
"""
import os
import sys

from openpyxl import Workbook

from builder import sheets as SH

OUT = "Master Industry Database.xlsx"

SHEET_PLAN = [
    ("", "Cover", None),
    ("", "Contents", None),
    ("1", "Company List", SH.company_list),
    ("2", "Production Capacity", SH.production_capacity),
    ("3", "Production Volume", SH.production_volume),
    ("4", "Revenue", SH.revenue),
    ("5", "EBITDA", SH.ebitda),
    ("6", "EBITDA per Ton", SH.ebitda_per_ton),
    ("7", "Steel Prices", SH.steel_prices),
    ("8", "Iron Ore Prices", SH.iron_ore_prices),
    ("9", "Coking Coal Prices", SH.coking_coal_prices),
    ("10", "Demand & Consumption", SH.demand),
    ("11", "Imports & Exports", SH.trade),
    ("12", "Government Policies", SH.policies),
    ("13", "Industry KPIs", SH.kpis),
    ("14", "Sources", SH.sources_sheet),
    ("15", "Data Dictionary", SH.dictionary_sheet),
    ("16", "Conflicts Log", SH.conflicts_sheet),
]

CONTENTS = [
    ("1", "Company List",
     "Coverage universe control record: 15 companies with identifiers, ownership, steelmaking "
     "route, assets, consolidation basis, currency and latest reported period.", "Tier 1"),
    ("2", "Production Capacity",
     "India capacity FY2024-FY2026 against the 300 Mtpa National Steel Policy target; company "
     "capacity and pipeline with approval status; SAIL plant-wise capacity and utilisation.",
     "Tier 1"),
    ("3", "Production Volume",
     "India crude steel, finished steel, DRI and pig iron output; world and China comparatives; "
     "company production and sales volumes with the denominator stated.", "Tier 1"),
    ("4", "Revenue",
     "Consolidated revenue from operations FY2022-FY2026 for all 15 names, with issuer headline "
     "measures carried separately where they differ, and a live CAGR formula.", "Tier 1"),
    ("5", "EBITDA",
     "EBITDA on two bases - as reported by the issuer and standardised - with the other-income "
     "treatment documented per company, plus a live margin formula.", "Tier 1"),
    ("6", "EBITDA per Ton",
     "The sector's core profitability metric, annual and quarterly, with the volume denominator "
     "stated on every row and FY2027 guidance where given.", "Tier 1"),
    ("7", "Steel Prices",
     "Domestic and international HRC and rebar observations with assessment basis and 52-week "
     "ranges; the FY2026 price path around the safeguard duty.", "Tier 1"),
    ("8", "Iron Ore Prices",
     "62% Fe CFR China benchmark: FY2016-FY2026 fiscal-year averages, calendar averages, monthly "
     "and quarterly; NMDC notified ex-mine prices; domestic and pellet assessments.", "Tier 1"),
    ("9", "Coking Coal Prices",
     "Premium HCC FOB Australia from both available public series, presented separately and never "
     "blended, plus an explicitly labelled thermal coal energy reference.", "Tier 1"),
    ("10", "Demand & Consumption",
     "Twelve fiscal years of apparent consumption, the full supply-demand balance that derives it, "
     "per capita intensity against world and China, and the latest quarter.", "Tier 1"),
    ("11", "Imports & Exports",
     "Twelve fiscal years of trade by volume and value; product mix; source and destination country "
     "shares; monthly net-trade turning points around the safeguard duty.", "Tier 1"),
    ("12", "Government Policies",
     "20 dated and scoped measures across trade remedies, standards, industrial policy, fiscal, "
     "decarbonisation, state support and overseas regimes, with assessed sector impact.", "Tier 1"),
    ("13", "Industry KPIs",
     "24 dashboard-ready indicators with stable KPI IDs spanning supply, demand, trade, input "
     "costs, realisations, policy and coverage.", "Tier 1"),
    ("14", "Sources",
     "47 sources with publisher, publication, date, section or page, period covered, URL and an "
     "explicit statement of what verification was actually performed.", "Reference"),
    ("15", "Data Dictionary",
     "Every column defined, with the reason it exists. Global metadata conventions plus "
     "sheet-specific definitions.", "Reference"),
    ("16", "Conflicts Log",
     "16 conflicts and definitional traps with competing figures, cause, resolution, figure carried "
     "forward and residual risk.", "Reference"),
]


def main():
    wb = Workbook()
    wb.remove(wb.active)
    made = {}
    for _, name, _ in SHEET_PLAN:
        made[name] = wb.create_sheet(title=name)

    SH.cover(made["Cover"])
    SH.contents(made["Contents"], CONTENTS)
    for num, name, fn in SHEET_PLAN:
        if fn is None:
            continue
        fn(made[name])

    wb.properties.title = "Master Industry Database - Indian Steel Industry"
    wb.properties.subject = "Indian steel industry institutional database"
    wb.properties.category = "Global Metals & Mining Research"
    wb.properties.keywords = ("Indian steel; JPC; Ministry of Steel; safeguard duty; "
                              "EBITDA per tonne; coking coal; iron ore; FY2026")
    wb.properties.description = (
        "Institutional database for the Indian steel industry. Data cutoff 03 August 2026. "
        "FY2026 = 1 April 2025 to 31 March 2026. Public sources only.")
    made["Cover"].sheet_view.tabSelected = True
    wb.active = 0
    wb.save(OUT)
    size = os.path.getsize(OUT)
    print("WROTE %s  (%.1f KB, %d sheets)" % (OUT, size / 1024.0, len(wb.sheetnames)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
