#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""QC for the copy-paste-safe Industry Financial Model.xlsx.

Includes a real recursive-descent Excel formula EVALUATOR, so the sheet's formulas are
actually computed and compared against fmodel/chain.py for all four scenarios. The
previous version of this workbook shipped broken; structural checks alone are not enough.
"""
import re
import sys

import openpyxl

MODEL = "Industry Financial Model.xlsx"
DB = "Master Industry Database.xlsx"
FAIL, WARN, OKC = [], [], []


def ok(c, msg):
    (OKC if c else FAIL).append(msg)
    return c


# ======================================================================================
# Minimal Excel formula evaluator (recursive descent, lazy IF/IFERROR/AND/OR)
# ======================================================================================
class Err(Exception):
    pass


def nf(v):
    """Coerce a single Excel value to float. Text in arithmetic is an error, as in Excel."""
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if v is None:
        return 0.0
    if isinstance(v, str):
        raise Err("text in arithmetic: %r" % v)
    return float(v)


def bc(a, b, f):
    """Broadcast a binary op over scalars and/or equal-length arrays, as Excel does."""
    la, lb = isinstance(a, list), isinstance(b, list)
    if not la and not lb:
        return f(a, b)
    n = max(len(a) if la else 1, len(b) if lb else 1)
    out = []
    for i in range(n):
        out.append(f(a[i] if la else a, b[i] if lb else b))
    return out


TOK = re.compile(r"""
    (?P<num>\d+\.?\d*(?:[eE][+-]?\d+)?)
  | (?P<str>"(?:[^"]|"")*")
  | (?P<rng>\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+)
  | (?P<cell>\$?[A-Z]{1,3}\$?\d+)
  | (?P<name>[A-Z][A-Z0-9_.]*)
  | (?P<op><>|<=|>=|[-+*/^()<>=,&])
  | (?P<ws>\s+)
""", re.X)


def lex(s):
    out, i = [], 0
    while i < len(s):
        m = TOK.match(s, i)
        if not m:
            raise Err("lex error at %d in %r" % (i, s))
        i = m.end()
        k = m.lastgroup
        if k == "ws":
            continue
        out.append((k, m.group()))
    return out


class P:
    def __init__(self, toks, ev):
        self.t, self.i, self.ev = toks, 0, ev

    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else (None, None)

    def eat(self, val=None):
        k, v = self.peek()
        if val and v != val:
            raise Err("expected %r got %r" % (val, v))
        self.i += 1
        return v

    def expr(self):
        left = self.add()
        while self.peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self.eat()
            right = self.add()
            a, b = self.ev.val(left), self.ev.val(right)

            def cmp(x, y, _op=op):
                if isinstance(x, str) or isinstance(y, str):
                    x, y = str(x), str(y)
                return {"=": x == y, "<>": x != y, "<": x < y, ">": x > y,
                        "<=": x <= y, ">=": x >= y}[_op]
            left = bc(a, b, cmp)
        return left

    def add(self):
        v = self.mul()
        while self.peek()[1] in ("+", "-", "&"):
            op = self.eat()
            r = self.mul()
            a, b = self.ev.val(v), self.ev.val(r)
            if op == "&":
                v = str(a) + str(b)
            elif op == "+":
                v = bc(a, b, lambda x, y: nf(x) + nf(y))
            else:
                v = bc(a, b, lambda x, y: nf(x) - nf(y))
        return v

    def mul(self):
        v = self.pw()
        while self.peek()[1] in ("*", "/"):
            op = self.eat()
            a, b = self.ev.val(v), self.ev.val(self.pw())
            if op == "*":
                v = bc(a, b, lambda x, y: nf(x) * nf(y))
            else:
                def dv(x, y):
                    if nf(y) == 0:
                        raise Err("div0")
                    return nf(x) / nf(y)
                v = bc(a, b, dv)
        return v

    def pw(self):
        v = self.un()
        while self.peek()[1] == "^":
            self.eat()
            a, b = self.ev.val(v), self.ev.val(self.un())
            v = bc(a, b, lambda x, y: nf(x) ** nf(y))
        return v

    def un(self):
        if self.peek()[1] in ("-", "+"):
            op = self.eat()
            v = self.ev.val(self.un())
            if op == "-":
                return bc(v, 0, lambda x, _y: -nf(x))
            return bc(v, 0, lambda x, _y: nf(x))
        return self.prim()

    def prim(self):
        k, v = self.peek()
        if v == "(":
            self.eat()
            e = self.expr()
            self.eat(")")
            return e
        if k == "num":
            self.eat()
            return float(v)
        if k == "str":
            self.eat()
            return v[1:-1].replace('""', '"')
        if k == "rng":
            self.eat()
            return ("RANGE", v.replace("$", ""))
        if k == "name":
            name = self.eat()
            if self.peek()[1] != "(":
                raise Err("unknown name %s" % name)
            self.eat("(")
            args, depth = [], 0
            start = self.i
            while True:
                kk, vv = self.peek()
                if vv is None:
                    raise Err("unterminated call")
                if vv == "(":
                    depth += 1
                elif vv == ")":
                    if depth == 0:
                        args.append(self.t[start:self.i])
                        self.eat()
                        break
                    depth -= 1
                elif vv == "," and depth == 0:
                    args.append(self.t[start:self.i])
                    self.eat()
                    start = self.i
                    continue
                self.i += 1
            return self.ev.call(name, args)
        if k == "cell":
            self.eat()
            return ("CELL", v.replace("$", ""))
        raise Err("unexpected %r" % v)


class Ev:
    def __init__(self, ws, overrides=None):
        self.ws = ws
        self.cache = {}
        self.ov = overrides or {}

    def raw(self, coord):
        if coord in self.ov:
            return self.ov[coord]
        return self.ws[coord].value

    def cell(self, coord):
        if coord in self.cache:
            v = self.cache[coord]
            if v == "__CYC__":
                raise Err("circular at %s" % coord)
            return v
        self.cache[coord] = "__CYC__"
        v = self.raw(coord)
        if isinstance(v, str) and v.startswith("="):
            v = self.run(v[1:])
        elif v is None:
            v = 0.0
        self.cache[coord] = v
        return v

    def run(self, body):
        p = P(lex(body), self)
        return self.val(p.expr())

    def val(self, v):
        if isinstance(v, tuple):
            if v[0] == "CELL":
                return self.cell(v[1])
            if v[0] == "RANGE":
                return self.rng(v[1])
        return v

    def num(self, v):
        v = self.val(v)
        if isinstance(v, bool):
            return 1.0 if v else 0.0
        if isinstance(v, str):
            raise Err("text in arithmetic: %r" % v)
        return float(v)

    def rng(self, ref):
        a, b = ref.split(":")
        c1, r1 = re.match(r"([A-Z]+)(\d+)", a).groups()
        c2, r2 = re.match(r"([A-Z]+)(\d+)", b).groups()
        from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
        out = []
        for rr in range(int(r1), int(r2) + 1):
            for cc in range(ci(c1), ci(c2) + 1):
                out.append(self.cell("%s%d" % (gl(cc), rr)))
        return out

    def sub(self, toks):
        return P(toks, self).expr()

    def call(self, name, args):
        A = args

        def E(i):
            return self.val(self.sub(A[i]))

        def N(i):
            return self.num(self.sub(A[i]))

        def flat(i):
            v = self.val(self.sub(A[i]))
            return v if isinstance(v, list) else [v]

        def nums(i):
            return [x for x in flat(i) if isinstance(x, (int, float))
                    and not isinstance(x, bool)]

        if name == "IF":
            c = E(0)
            c = bool(c) if not isinstance(c, str) else (c != "")
            if c:
                return E(1)
            return E(2) if len(A) > 2 else False
        if name == "IFERROR":
            try:
                return E(0)
            except (Err, ZeroDivisionError, ValueError, TypeError):
                return E(1)
        if name == "AND":
            for i in range(len(A)):
                v = E(i)
                if not (bool(v) if not isinstance(v, str) else v != ""):
                    return False
            return True
        if name == "OR":
            for i in range(len(A)):
                if bool(E(i)):
                    return True
            return False
        if name == "NOT":
            return not bool(E(0))
        if name == "CHOOSE":
            k = int(round(N(0)))
            return E(k)
        if name == "INDEX":
            arr = flat(0)
            k = int(round(N(1)))
            if k < 1 or k > len(arr):
                raise Err("INDEX out of range")
            return arr[k - 1]
        if name == "MIN":
            v = []
            for i in range(len(A)):
                v += nums(i)
            return min(v) if v else 0.0
        if name == "MAX":
            v = []
            for i in range(len(A)):
                v += nums(i)
            return max(v) if v else 0.0
        if name == "SUM":
            t = 0.0
            for i in range(len(A)):
                t += sum(nums(i))
            return t
        if name == "ABS":
            v = self.val(self.sub(A[0]))
            return bc(v, 0, lambda x, _y: abs(nf(x)))
        if name == "COUNT":
            t = 0
            for i in range(len(A)):
                t += len(nums(i))
            return float(t)
        if name == "MEDIAN":
            v = []
            for i in range(len(A)):
                v += nums(i)
            if not v:
                raise Err("MEDIAN empty")
            v.sort()
            n = len(v)
            return v[n // 2] if n % 2 else (v[n // 2 - 1] + v[n // 2]) / 2.0
        if name == "ISNUMBER":
            try:
                v = E(0)
            except Err:
                return False
            return isinstance(v, (int, float)) and not isinstance(v, bool)
        if name == "ISBLANK":
            tok = A[0]
            if len(tok) == 1 and tok[0][0] == "cell":
                return self.raw(tok[0][1].replace("$", "")) is None
            return False
        if name == "COUNTIF":
            rng = flat(0)
            crit = E(1)
            if isinstance(crit, str) and crit.startswith("*") and crit.endswith("*"):
                sub = crit.strip("*")
                return float(sum(1 for x in rng if isinstance(x, str) and sub in x))
            return float(sum(1 for x in rng if x == crit))
        if name == "SUMIFS":
            sm = flat(0)
            out = [True] * len(sm)
            i = 1
            while i + 1 < len(A) + 1 and i + 1 <= len(A):
                cr = flat(i)
                cv = E(i + 1)
                for j in range(len(sm)):
                    if j >= len(cr) or cr[j] != cv:
                        out[j] = False
                i += 2
            return float(sum(sm[j] for j in range(len(sm))
                             if out[j] and isinstance(sm[j], (int, float))))
        if name == "SUMPRODUCT":
            arrs = [flat(i) for i in range(len(A))]
            n = max(len(a) for a in arrs)
            t = 0.0
            for j in range(n):
                p = 1.0
                for a in arrs:
                    x = a[j] if j < len(a) else 0
                    if isinstance(x, bool):
                        x = 1.0 if x else 0.0
                    if isinstance(x, str):
                        x = 0.0
                    p *= float(x)
                t += p
            return t
        raise Err("unsupported function %s" % name)


# ======================================================================================
wb = openpyxl.load_workbook(MODEL)
print("SHEETS (%d): %s\n" % (len(wb.sheetnames), " | ".join(wb.sheetnames)))

# ---- 1. copy-paste safety (the whole point of this rebuild)
print("=== COPY-PASTE SAFETY ===")
xref = []
nform = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                nform += 1
                if "!" in v:
                    xref.append("%s!%s : %s" % (ws.title, c.coordinate, v[:70]))
print("Formulas: %d" % nform)
ok(not xref, "no cross-sheet references in any formula")
print("Cross-sheet references            : %d %s" % (len(xref), "OK" if not xref else "FAIL"))
for x in xref[:5]:
    print("    ", x)
dn = list(wb.defined_names)
ok(not dn, "no defined names")
print("Defined names                     : %d %s" % (len(dn), "OK" if not dn else "FAIL"))
tbl = sum(len(ws.tables or {}) for ws in wb.worksheets)
ok(tbl == 0, "no Excel Tables")
print("Excel Tables                      : %d %s" % (tbl, "OK" if tbl == 0 else "FAIL"))
mg = sum(len(ws.merged_cells.ranges) for ws in wb.worksheets)
ok(mg == 0, "no merged cells")
print("Merged cells                      : %d %s" % (mg, "OK" if mg == 0 else "FAIL"))

# text that should have been a formula
textref = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not isinstance(v, str) or v.startswith("="):
                continue
            t = v.strip()
            # Only a defect in a NUMERIC DATA column (D..L). Columns A-C hold section and
            # driver IDs ("S0", "D01", "K12") and columns M+ hold Source IDs ("S01", "S17"),
            # all of which legitimately look like cell references but are labels.
            if c.column < 4 or c.column > 12:
                continue
            if "!" in t or re.match(r"^\$?[A-Z]{1,3}\$?\d+$", t):
                textref.append("%s!%s = %r" % (ws.title, c.coordinate, v))
ok(not textref, "no cell reference left as literal text")
print("Text-that-should-be-formula       : %d %s"
      % (len(textref), "OK" if not textref else "FAIL"))
for t in textref[:5]:
    print("    ", t)

# ---- 2. evaluate the Model sheet and compare to chain.py
print("\n=== FORMULA EVALUATION vs fmodel/chain.py ===")
sys.path.insert(0, ".")
from fmodel import chain as C                                              # noqa: E402
from fmodel.assumptions import SCENARIOS                                   # noqa: E402

ms = wb["Model"]
# locate the scenario cell and the key rows by their column-B labels
# Map label -> row. Take the LAST occurrence and SKIP the S1 driver-matrix rows, which are
# identified by column C holding a scenario name rather than a unit. Several driver names
# ("Finished steel imports") also appear as calculation lines further down; the calculation
# line is the one to test.
lab = {}
for r in range(1, ms.max_row + 1):
    v = ms.cell(r, 2).value
    cu = ms.cell(r, 3).value
    if isinstance(cu, str) and cu.strip() in SCENARIOS:
        continue
    if isinstance(v, str) and v.strip():
        lab[v.strip()] = r
scen_row = None
for r in range(1, 60):
    if isinstance(ms.cell(r, 2).value, str) and "SCENARIO (enter" in ms.cell(r, 2).value:
        scen_row = r
        break
ok(scen_row is not None, "scenario cell located")
SCELL = "D%d" % scen_row
print("Scenario cell: %s" % SCELL)

CHECKS = [
    ("Apparent finished steel consumption", "cons", 0.01),
    ("Crude steel production", "crude", 0.01),
    ("Finished steel production", "prod", 0.01),
    ("Crude steel capacity (closing)", "cap", 0.01),
    ("Capacity utilisation", "util", 0.0001),
    ("Blended realisation - EFFECTIVE", "realn", 1.0),
    ("Industry cash cost per tonne", "cost", 1.0),
    ("INDUSTRY EBITDA PER TONNE", "ebt", 1.0),
    ("INDUSTRY REVENUE", "rev", 5.0),
    ("INDUSTRY EBITDA", "ebitda", 5.0),
    ("INDUSTRY EBITDA MARGIN", "margin", 0.0005),
    ("UNLEVERED FREE CASH FLOW", "fcf", 20.0),
    ("Per capita consumption", "percap", 0.05),
    ("Finished steel imports", "imports", 0.01),
    ("Finished steel exports", "exports", 0.01),
]
FCOLS = ["E", "F", "G", "H", "I", "J", "K"]
allok = True
for si, sname in enumerate(SCENARIOS, start=1):
    ev = Ev(ms, overrides={SCELL: si})
    bad = 0
    worst = ("", 0.0)
    for label, key, tol in CHECKS:
        if label not in lab:
            FAIL.append("label not found on Model sheet: %s" % label)
            continue
        rr = lab[label]
        for i, col in enumerate(FCOLS):
            try:
                got = ev.cell("%s%d" % (col, rr))
            except Err as e:
                FAIL.append("%s / %s %s: eval error %s" % (sname, label, col, e))
                bad += 1
                continue
            exp = C.ALL[sname][key][i]
            d = abs(float(got) - float(exp))
            if d > tol:
                bad += 1
                if d > worst[1]:
                    worst = ("%s %s(%s) got %.4f exp %.4f" % (sname, label, col, got, exp), d)
    status = "OK" if bad == 0 else "FAIL (%d cells)" % bad
    print("  %-13s %d line items x 7 years = %3d cells   %s"
          % (sname, len(CHECKS), len(CHECKS) * 7, status))
    if bad:
        allok = False
        print("      worst: %s" % worst[0])
ok(allok, "Model sheet formulas reproduce chain.py for all four scenarios")

# ---- 3. audit-check formulas evaluate to PASS
print("\n=== IN-WORKBOOK AUDIT CHECKS (evaluated) ===")
ev = Ev(ms, overrides={SCELL: 1})
res = {}
for r in range(1, ms.max_row + 1):
    a = ms.cell(r, 1).value
    if isinstance(a, str) and re.match(r"^K\d\d$", a):
        try:
            res[a] = ev.cell("D%d" % r)
        except Err as e:
            res[a] = "EVAL ERROR %s" % e
chk_fail = [k for k, v in res.items() if isinstance(v, str) and "FAIL" in v]
chk_warn = [k for k, v in res.items() if isinstance(v, str) and "WARN" in v]
chk_err = [k for k, v in res.items() if isinstance(v, str) and "ERROR" in v]
print("Checks evaluated: %d   PASS: %d   WARN: %d   FAIL: %d   ERROR: %d"
      % (len(res), sum(1 for v in res.values() if isinstance(v, str) and v.startswith("PASS")),
         len(chk_warn), len(chk_fail), len(chk_err)))
for k in sorted(chk_warn):
    print("   WARN  %s -> %r" % (k, res[k]))
for k in chk_fail + chk_err:
    print("   %s -> %r" % (k, res[k]))
ok(not chk_fail and not chk_err, "all in-workbook audit checks pass in the Base Case")

# ---- 4. reconcile to the database
print("\n=== RECONCILIATION TO MASTER INDUSTRY DATABASE ===")
try:
    dbw = openpyxl.load_workbook(DB, data_only=True)
    pv = dbw["Production Volume"]
    got = None
    for r in range(1, pv.max_row + 1):
        if pv.cell(r, 1).value == "India crude steel production":
            got = pv.cell(r, 7).value
    pc = dbw["Production Capacity"]
    gcap = None
    for r in range(1, pc.max_row + 1):
        if pc.cell(r, 1).value == "India crude steel capacity":
            gcap = pc.cell(r, 5).value
    dc = dbw["Demand & Consumption"]
    gc = gp = None
    for r in range(1, dc.max_row + 1):
        v = dc.cell(r, 1).value
        if isinstance(v, str):
            if v.startswith("7. Apparent"):
                gc = dc.cell(r, 6).value
            if v.startswith("2. Finished steel production"):
                gp = dc.cell(r, 6).value
    for nm, g, e in [("crude steel production FY2026", got, C.CRUDE26),
                     ("crude steel capacity FY2026", gcap, C.CAP26),
                     ("consumption FY2026", gc, C.CONS26),
                     ("finished production FY2026", gp, C.PROD26)]:
        good = isinstance(g, (int, float)) and abs(g - e) < 0.01
        print("  %-34s DB=%-9s model=%-9s %s" % (nm, g, e, "OK" if good else "MISMATCH"))
        ok(good, "DB reconciliation: %s" % nm)
except Exception as exc:                                                   # noqa: BLE001
    WARN.append("database not re-checked: %s" % exc)
    print("  skipped: %s" % exc)

# ---- 5. file integrity
print("\n=== FILE INTEGRITY ===")
import zipfile                                                            # noqa: E402
z = zipfile.ZipFile(MODEL)
bad = z.testzip()
ok(bad is None, "zip container intact")
print("Zip container: %s   parts: %d" % ("OK" if bad is None else bad, len(z.namelist())))
wb2 = openpyxl.load_workbook(MODEL)
ok(len(wb2.sheetnames) == len(wb.sheetnames), "round-trips through openpyxl")
print("Round-trip reload: OK (%d sheets)" % len(wb2.sheetnames))

print()
if WARN:
    print("WARNINGS (%d):" % len(WARN))
    for w in WARN:
        print("   - %s" % w)
if FAIL:
    print("FAILURES (%d):" % len(FAIL))
    for f in FAIL[:40]:
        print("   - %s" % f)
    sys.exit(1)
print("QC PASSED - %d checks OK, %d warnings, 0 failures." % (len(OKC), len(WARN)))
