#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Quality-control audit for Master Industry Database.xlsx."""
import re
import sys

import openpyxl

PATH = "Master Industry Database.xlsx"
FAIL = []
WARN = []


def check(cond, msg):
    if not cond:
        FAIL.append(msg)


def warn(cond, msg):
    if not cond:
        WARN.append(msg)


wb = openpyxl.load_workbook(PATH)
print("SHEETS (%d): %s\n" % (len(wb.sheetnames), " | ".join(wb.sheetnames)))

REQUIRED = ["Company List", "Production Capacity", "Production Volume", "Revenue", "EBITDA",
            "EBITDA per Ton", "Steel Prices", "Iron Ore Prices", "Coking Coal Prices",
            "Demand & Consumption", "Imports & Exports", "Government Policies",
            "Industry KPIs", "Sources"]
for s in REQUIRED:
    check(s in wb.sheetnames, "MISSING mandated worksheet: %s" % s)

all_tables = {}
total_cells = 0
na_cells = 0
formula_cells = 0
merged_total = 0

for ws in wb.worksheets:
    merged_total += len(ws.merged_cells.ranges)
    check(not ws.merged_cells.ranges,
          "%s: has %d merged range(s); house style forbids merged cells"
          % (ws.title, len(ws.merged_cells.ranges)))
    check(ws.freeze_panes is not None or ws.title in ("Cover",),
          "%s: no freeze panes" % ws.title)
    check(ws.sheet_view.showGridLines is False, "%s: gridlines still visible" % ws.title)

    for tname, tref in (ws.tables or {}).items():
        check(tname not in all_tables,
              "DUPLICATE table name %s (in %s and %s)" % (tname, ws.title,
                                                          all_tables.get(tname)))
        all_tables[tname] = ws.title
        m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", tref)
        r1, r2 = int(m.group(2)), int(m.group(4))
        check(r2 > r1, "%s/%s: table has no data rows (%s)" % (ws.title, tname, tref))
        # header row must be fully populated and unique
        c1 = openpyxl.utils.column_index_from_string(m.group(1))
        c2 = openpyxl.utils.column_index_from_string(m.group(3))
        hdrs = [ws.cell(r1, c).value for c in range(c1, c2 + 1)]
        check(all(h not in (None, "") for h in hdrs),
              "%s/%s: blank header cell in table range" % (ws.title, tname))
        check(len(set(hdrs)) == len(hdrs),
              "%s/%s: duplicate header names %s" % (ws.title, tname,
                                                    [h for h in hdrs if hdrs.count(h) > 1]))

    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            total_cells += 1
            if isinstance(c.value, str):
                if c.value.startswith("="):
                    formula_cells += 1
                    check("#REF" not in c.value and "#VALUE" not in c.value,
                          "%s!%s: formula contains an error token" % (ws.title, c.coordinate))
                elif c.value == "Data Not Publicly Available":
                    na_cells += 1
            elif isinstance(c.value, (int, float)):
                check(c.number_format not in ("General",),
                      "%s!%s: numeric cell with General number format (value=%s)"
                      % (ws.title, c.coordinate, c.value))

print("Tables: %d  |  populated cells: %d  |  formulas: %d  |  "
      "'Data Not Publicly Available': %d  |  merged ranges: %d\n"
      % (len(all_tables), total_cells, formula_cells, na_cells, merged_total))

# ---------------------------------------------------------------- value spot-checks
def find_row(ws, col, needle, hdr_needle=None):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and v.strip() == needle:
            return r
    return None


SPOT = []
ws = wb["Revenue"]
# locate header row of tbl_Revenue
for r in range(1, 60):
    if ws.cell(r, 1).value == "Company":
        hdr = r
        break
cols = {ws.cell(hdr, c).value: c for c in range(1, 20)}
for company, fy26 in [("Tata Steel", 232140), ("JSW Steel", 185470), ("SAIL", 110811),
                      ("Jindal Steel", 53225), ("Jindal Stainless", 42955),
                      ("Shyam Metalics", 18552)]:
    rr = find_row(ws, 1, company)
    got = ws.cell(rr, cols["FY2026"]).value if rr else None
    SPOT.append(("Revenue FY2026 %s" % company, fy26, got))

ws = wb["EBITDA"]
for r in range(1, 60):
    if ws.cell(r, 1).value == "Company":
        hdr = r
        break
cols = {ws.cell(hdr, c).value: c for c in range(1, 25)}
for company, rep, std in [("Tata Steel", 34848, 34352), ("JSW Steel", 29821, 29464),
                          ("SAIL", 13146, 12000), ("Jindal Stainless", 5560, 5560),
                          ("Shyam Metalics", 2537, 2333)]:
    rr = find_row(ws, 1, company)
    SPOT.append(("EBITDA reported FY2026 %s" % company, rep,
                 ws.cell(rr, cols["Reported FY2026"]).value))
    SPOT.append(("EBITDA standardised FY2026 %s" % company, std,
                 ws.cell(rr, cols["Standardised FY2026"]).value))

ws = wb["Demand & Consumption"]
rr = find_row(ws, 1, "FY2026")
SPOT.append(("Consumption FY2026 (Mt)", 164.19, ws.cell(rr, 2).value))
SPOT.append(("Crude steel FY2026 (Mt)", 168.42, ws.cell(rr, 5).value))
SPOT.append(("Capacity FY2026 (Mtpa)", 220.4, ws.cell(rr, 6).value))

ws = wb["Iron Ore Prices"]
rr = find_row(ws, 1, "FY2026")
SPOT.append(("Iron ore FY2026 avg (US$/dmt)", 100.52, ws.cell(rr, 2).value))

ws = wb["Imports & Exports"]
rr = find_row(ws, 1, "FY2026")
SPOT.append(("Imports FY2026 (Mt)", 6.524, ws.cell(rr, 2).value))
SPOT.append(("Exports FY2026 (Mt)", 6.602, ws.cell(rr, 3).value))

bad = 0
print("VALUE SPOT-CHECKS")
for label, expect, got in SPOT:
    ok = (isinstance(got, (int, float)) and abs(got - expect) < 0.005)
    if not ok:
        bad += 1
    print("  %-46s expect %-12s got %-12s %s" % (label, expect, got, "OK" if ok else "<< FAIL"))
check(bad == 0, "%d value spot-check(s) failed" % bad)

# ---------------------------------------------------------------- internal consistency
ws = wb["Imports & Exports"]
rr = find_row(ws, 1, "FY2026")
f = ws.cell(rr, 4).value
check(isinstance(f, str) and f.startswith("="), "Net Trade column is not a live formula")

ws = wb["Revenue"]
rr = find_row(ws, 1, "Tata Steel")
f = ws.cell(rr, 13).value
check(isinstance(f, str) and f.startswith("="), "Revenue CAGR column is not a live formula")

print()
if WARN:
    print("WARNINGS (%d):" % len(WARN))
    for w in WARN:
        print("  - %s" % w)
if FAIL:
    print("FAILURES (%d):" % len(FAIL))
    for f in FAIL:
        print("  - %s" % f)
    sys.exit(1)
print("QC PASSED - no failures.")
